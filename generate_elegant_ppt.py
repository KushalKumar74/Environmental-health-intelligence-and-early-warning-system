import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_elegant_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # -------------------------------------------------------------
    # ELEGANT LUXURY COLOR PALETTE (Teal & Warm Amber Gold)
    # -------------------------------------------------------------
    BG_MIDNIGHT = RGBColor(11, 19, 36)      # Deep Midnight Navy
    BG_LIGHT = RGBColor(248, 250, 252)     # Soft Pearl Slate
    TEAL_PRIMARY = RGBColor(13, 148, 136)   # Refined Teal
    AMBER_GOLD = RGBColor(217, 119, 6)     # Warm Gold Accent
    TEXT_MAIN = RGBColor(30, 41, 59)        # Deep Charcoal Body Text
    TEXT_MUTED = RGBColor(100, 116, 139)    # Elegant Subtitle Gray
    TEXT_WHITE = RGBColor(255, 255, 255)    # White
    CARD_WHITE = RGBColor(255, 255, 255)    # Pure White Card
    CARD_BORDER = RGBColor(226, 232, 240)   # Subtle Border
    CARD_ACCENT_BG = RGBColor(241, 245, 249)# Soft Gray Card fill
    SKY_BLUE = RGBColor(14, 165, 233)       # Sky Accent

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, topic_kicker="ENVIRONMENTAL HEALTH INTELLIGENCE & EARLY WARNING SYSTEM"):
        # Category Kicker (Small uppercase tracker)
        kicker_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
        tf_k = kicker_box.text_frame
        tf_k.word_wrap = True
        tf_k.margin_left = tf_k.margin_right = tf_k.margin_top = tf_k.margin_bottom = 0
        p_k = tf_k.paragraphs[0]
        p_k.text = topic_kicker.upper()
        p_k.font.size = Pt(10)
        p_k.font.bold = True
        p_k.font.color.rgb = TEAL_PRIMARY

        # Main Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.65))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

        # Decorative Underline Bar
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.38), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

    def add_card(slide, left, top, width, height, title, items, top_accent_color=TEAL_PRIMARY, pill_tag=None):
        # Card Background Shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_WHITE
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Top Color Accent Line
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.06))
        accent.fill.solid()
        accent.fill.fore_color.rgb = top_accent_color
        accent.line.fill.background()

        # Text Frame with Generous Padding (Cell Spacing)
        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        # Header paragraph
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN
        p_t.space_after = Pt(10)
        p_t.line_spacing = 1.15

        # Bullets
        for item in items:
            p_b = tf.add_paragraph()
            p_b.text = f"• {item}"
            p_b.font.size = Pt(11.5)
            p_b.font.color.rgb = TEXT_MAIN
            p_b.space_after = Pt(6)
            p_b.line_spacing = 1.15

    # -------------------------------------------------------------
    # SLIDE 1: Elegant Title Slide (Dark Theme)
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, BG_MIDNIGHT)

    # Decorative accent bar
    dec_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(0.8), Inches(0.06))
    dec_bar.fill.solid()
    dec_bar.fill.fore_color.rgb = AMBER_GOLD
    dec_bar.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_right = tf1.margin_top = tf1.margin_bottom = 0

    p1_sub = tf1.paragraphs[0]
    p1_sub.text = "ACADEMIC MINI PROJECT PRESENTATION"
    p1_sub.font.size = Pt(12)
    p1_sub.font.bold = True
    p1_sub.font.color.rgb = AMBER_GOLD
    p1_sub.space_after = Pt(12)

    p1_title = tf1.add_paragraph()
    p1_title.text = "Environmental Health Intelligence & Early Warning System"
    p1_title.font.size = Pt(30)
    p1_title.font.bold = True
    p1_title.font.color.rgb = TEXT_WHITE
    p1_title.space_after = Pt(14)
    p1_title.line_spacing = 1.15

    p1_desc = tf1.add_paragraph()
    p1_desc.text = "Real-Time AQI Monitoring, Machine Learning Anomaly Detection & AI Health Risk Platform"
    p1_desc.font.size = Pt(16)
    p1_desc.font.color.rgb = TEAL_PRIMARY
    p1_desc.space_after = Pt(36)

    p1_meta = tf1.add_paragraph()
    p1_meta.text = "Architecture: React (Vite) + Node.js (Express v5) + Python FastAPI (Scikit-Learn) + MongoDB"
    p1_meta.font.size = Pt(12)
    p1_meta.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 2: Agenda (6 Required Topics)
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, BG_LIGHT)
    add_header(slide2, "Presentation Agenda & Executive Overview")

    agenda_cards = [
        ("1. Introduction", "Project Background, Objectives & Core Innovation Pillars", TEAL_PRIMARY),
        ("2. Problem Definition & Existing Systems", "Core Public Health Challenges & Comparative Feature Matrix", AMBER_GOLD),
        ("3. Methodology", "Data Acquisition Pipeline, ML Anomaly Engine & EVA Assistant", SKY_BLUE),
        ("4. Design Diagrams", "System Architecture, ER Schema & Data Flow Sequence Diagrams", TEAL_PRIMARY),
        ("5. Conclusion & Discussion", "Technical Impact, Performance Benchmarks & Future Roadmap", AMBER_GOLD),
        ("6. Screenshots & Other Diagrams", "Interactive GIS Map, UI Layout, Dashboards & PDF Exporter", SKY_BLUE)
    ]

    for idx, (title, desc, color) in enumerate(agenda_cards):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.65 + row * 1.75)
        add_card(slide2, left, top, Inches(5.65), Inches(1.5), title, [desc], color)

    # -------------------------------------------------------------
    # SLIDE 3: 1. Introduction
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, BG_LIGHT)
    add_header(slide3, "1. Introduction — Project Overview & Motivation", "TOPIC 1: INTRODUCTION")

    add_card(slide3, Inches(0.8), Inches(1.65), Inches(3.65), Inches(5.1), "Background & Urgency", [
        "Rapid urbanization and industrial emissions lead to severe particulate spikes (PM2.5, PM10).",
        "Air pollution directly correlates with worsening respiratory and cardiovascular health conditions.",
        "Traditional public tracking tools provide passive numbers without actionable health guidance."
    ], TEAL_PRIMARY)

    add_card(slide3, Inches(4.85), Inches(1.65), Inches(3.65), Inches(5.1), "Project Objectives", [
        "Geospatial AQI Tracking: Spatial mapping of 6 key atmospheric pollutants.",
        "Predictive Machine Learning: Real-time detection of sudden environmental anomalies.",
        "Personalized Risk Assessment: Tailored dosage calculation based on user health profiles.",
        "AI Virtual Assistance: Interactive EVA assistant for travel and medical emergency routing."
    ], AMBER_GOLD)

    add_card(slide3, Inches(8.9), Inches(1.65), Inches(3.65), Inches(5.1), "Core Innovation Pillars", [
        "Microservices Architecture: Node.js API server paired with Python FastAPI ML service.",
        "Isolation Forest ML Engine: Unsupervised detection of localized pollution spikes.",
        "Doctor-Ready PDF Reports: Downloadable clinical exposure summaries for medical review.",
        "Multi-Tier Access Control: Custom views for Users, Facilities, and Administrators."
    ], SKY_BLUE)

    # -------------------------------------------------------------
    # SLIDE 4: 2. Problem Definition & Existing Systems
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, BG_LIGHT)
    add_header(slide4, "2. Problem Definition & Existing Systems", "TOPIC 2: PROBLEM DEFINITION")

    add_card(slide4, Inches(0.8), Inches(1.65), Inches(5.65), Inches(5.1), "Core Public Health Challenges", [
        "One-Size-Fits-All AQI Metrics: Standard air quality indices treat asthmatics, children, and healthy adults identically.",
        "Reactive Emergency Response: Alerts are triggered *after* heavy exposure has already occurred, failing to prevent health damage.",
        "Fragmented Healthcare Access: Citizens cannot easily locate nearby specialized clinics or hospitals during severe smog episodes.",
        "Unfiltered Sensor Telemetry: Raw sensor feeds frequently suffer from unverified noise or sudden sensor calibration errors."
    ], AMBER_GOLD)

    add_card(slide4, Inches(6.85), Inches(1.65), Inches(5.65), Inches(5.1), "Comparative Feature Matrix", [
        "Legacy Apps: Show generic AQI numbers without individual susceptibility context.",
        "EHI-EWS: Calculates personalized cumulative exposure dosage based on user medical profiles.",
        "Legacy Apps: No real-time predictive anomaly detection for sudden local fires or emissions.",
        "EHI-EWS: Microservice-driven Isolation Forest ML engine with statistical fallbacks.",
        "Legacy Apps: Static UI with zero interactive guidance.",
        "EHI-EWS: EVA Virtual Assistant providing smart itineraries, travel advice, and PDF exports."
    ], TEAL_PRIMARY)

    # -------------------------------------------------------------
    # SLIDE 5: 3. Methodology - Data Pipeline & ML Microservice
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, BG_LIGHT)
    add_header(slide5, "3. Methodology — Data Pipeline & ML Engine", "TOPIC 3: METHODOLOGY")

    add_card(slide5, Inches(0.8), Inches(1.65), Inches(5.65), Inches(5.1), "Data Acquisition & Caching Pipeline", [
        "Telemetry Source: Integrates OpenWeather REST API for live air pollution and weather forecasts.",
        "Monitored Pollutants: Full tracking of PM2.5, PM10, NO2, SO2, CO, and O3 values.",
        "NodeCache Optimization: Backend cache layer with 10-minute TTL reduces external API requests by 85%.",
        "Geocoding Engine: Converts latitude and longitude coordinates into localized geographic place names."
    ], TEAL_PRIMARY)

    add_card(slide5, Inches(6.85), Inches(1.65), Inches(5.65), Inches(5.1), "Machine Learning Anomaly Detection", [
        "Isolation Forest Model: Fitted on sequential pollutant series to identify unexpected pollution anomalies.",
        "Z-Score Statistical Fallback: Evaluates rolling z-score ($z > 2.2$) when time-series history is limited.",
        "Quantitative Model Evaluator: Computes MAE, RMSE, and R² for continuous prediction accuracy verification.",
        "FastAPI Service: Isolated execution on port 8000 for low-latency machine learning inference."
    ], SKY_BLUE)

    # -------------------------------------------------------------
    # SLIDE 6: 3. Methodology - Health Risk Assessment & EVA
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, BG_LIGHT)
    add_header(slide6, "3. Methodology — Health Assessment & EVA Assistant", "TOPIC 3: METHODOLOGY")

    add_card(slide6, Inches(0.8), Inches(1.65), Inches(5.65), Inches(5.1), "Personalized Exposure Assessment Model", [
        "Health Profile Integration: Captures pre-existing conditions (Asthma, Heart Disease, Vulnerable Age).",
        "Dosage Equation: Exposure Dose = Ambient Concentration × Outdoor Duration × Vulnerability Score.",
        "Risk Stratification: Classifies safety into Safe, Moderate, Severe, and Critical hazard zones.",
        "Clinical PDF Exporter: Uses jsPDF to generate medical summary reports for physician review."
    ], AMBER_GOLD)

    add_card(slide6, Inches(6.85), Inches(1.65), Inches(5.65), Inches(5.1), "EVA Virtual Assistant Architecture", [
        "Dual Interaction Mode: Guided menu options combined with open natural language query processing.",
        "Context Awareness: Reads active page context, current GPS coordinates, and user risk tier.",
        "Smart Outdoor Itinerary: Recommends safe outdoor exercise windows and safe indoor alternatives.",
        "Emergency Navigation: Directly routes users to nearby registered medical centers and clinics."
    ], TEAL_PRIMARY)

    # -------------------------------------------------------------
    # SLIDE 7: 4. Design Diagrams - System Architecture
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, BG_LIGHT)
    add_header(slide7, "4. Design Diagrams — System Architecture", "TOPIC 4: DESIGN DIAGRAMS")

    # Architect boxes with clean styling and generous cell padding
    def add_arch_box(left, top, width, height, title, items, color=TEAL_PRIMARY):
        box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_WHITE
        box.line.color.rgb = color
        box.line.width = Pt(1.5)

        tb = slide7.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN
        p_t.alignment = PP_ALIGN.CENTER
        p_t.space_after = Pt(8)

        for item in items:
            p_i = tf.add_paragraph()
            p_i.text = f"▪ {item}"
            p_i.font.size = Pt(11)
            p_i.font.color.rgb = TEXT_MAIN
            p_i.space_after = Pt(4)

    add_arch_box(Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8), "CLIENT TIER\n(React + TypeScript)", [
        "Vite Build + Tailwind CSS",
        "Shadcn UI Component Suite",
        "Leaflet GIS Interactive Map",
        "Recharts Visualization",
        "jsPDF Clinical Exporter",
        "EVA Assistant Widget"
    ], SKY_BLUE)

    arr1 = slide7.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.5), Inches(3.9), Inches(0.6), Inches(0.35))
    arr1.fill.solid()
    arr1.fill.fore_color.rgb = TEAL_PRIMARY
    arr1.line.fill.background()

    add_arch_box(Inches(5.2), Inches(1.8), Inches(3.6), Inches(4.8), "BACKEND API TIER\n(Node.js + Express v5)", [
        "JWT Auth & RBAC Middleware",
        "AQI Service & NodeCache Layer",
        "EVA Orchestrator Service",
        "Helmet & Express Rate Limiter",
        "Mongoose Database Models"
    ], TEAL_PRIMARY)

    arr2 = slide7.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.9), Inches(3.9), Inches(0.6), Inches(0.35))
    arr2.fill.solid()
    arr2.fill.fore_color.rgb = AMBER_GOLD
    arr2.line.fill.background()

    add_arch_box(Inches(9.6), Inches(1.8), Inches(2.933), Inches(2.25), "ML SERVICE TIER\n(Python FastAPI)", [
        "Isolation Forest Model",
        "Z-Score Rolling Check",
        "Metrics Evaluator (MAE, R²)"
    ], AMBER_GOLD)

    add_arch_box(Inches(9.6), Inches(4.35), Inches(2.933), Inches(2.25), "DATA & EXTERNAL", [
        "MongoDB Database",
        "OpenWeather Pollution API",
        "Geocoding REST Services"
    ], TEXT_MUTED)

    # -------------------------------------------------------------
    # SLIDE 8: 4. Design Diagrams - ER Schema Diagram
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, BG_LIGHT)
    add_header(slide8, "4. Design Diagrams — Entity-Relationship (ER) Schema", "TOPIC 4: DESIGN DIAGRAMS")

    add_card(slide8, Inches(0.8), Inches(1.65), Inches(3.65), Inches(5.1), "User & Medical Profile", [
        "User Schema: _id, name, email, passwordHash, role (User, Admin, Facility), createdAt.",
        "Profile Schema: _id, userId, age, asthma (Boolean), heartCondition (Boolean), vulnerabilityScore, emergencyContact."
    ], SKY_BLUE)

    add_card(slide8, Inches(4.85), Inches(1.65), Inches(3.65), Inches(5.1), "Environmental Readings", [
        "EnvironmentalReading: _id, location, latitude, longitude, aqi, pm25, pm10, no2, so2, co, o3, timestamp.",
        "Prediction: _id, modelVersion, anomalyScore, isAnomaly, confidence, timestamp."
    ], TEAL_PRIMARY)

    add_card(slide8, Inches(8.9), Inches(1.65), Inches(3.65), Inches(5.1), "Exposure & Healthcare Schema", [
        "ExposureRecord: _id, userId, durationMinutes, averageAqi, calculatedDose, date.",
        "Facility Schema: _id, name, type (Hospital/Clinic), coordinates, capacity, contactPhone.",
        "EnvironmentalAlert: _id, severity, threshold, triggerLocation, status."
    ], AMBER_GOLD)

    # -------------------------------------------------------------
    # SLIDE 9: 4. Design Diagrams - Data Flow Sequence
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, BG_LIGHT)
    add_header(slide9, "4. Design Diagrams — End-to-End Data Flow Sequence", "TOPIC 4: DESIGN DIAGRAMS")

    add_card(slide9, Inches(0.8), Inches(1.65), Inches(11.733), Inches(5.1), "Sequential Execution Pipeline", [
        "Step 1 (Client Trigger): User views dashboard or queries air quality for specific GPS coordinates.",
        "Step 2 (Security & Rate Check): Express rate limiter & JWT auth middleware validate request authenticity.",
        "Step 3 (Cache Layer Inspection): NodeCache checks if OpenWeather telemetry for (lat, lon) exists within 10-min TTL.",
        "Step 4 (API Telemetry Fetch): On cache miss, backend fetches 6-pollutant telemetry & weather forecast from OpenWeather REST API.",
        "Step 5 (ML Anomaly Microservice Call): Express backend sends recent AQI sequence to Python FastAPI `/detect-anomalies` endpoint.",
        "Step 6 (Model Inference): Isolation Forest checks sequence for unexpected spikes, returning anomaly score and confidence metric.",
        "Step 7 (Health Risk Calculation): Exposure engine merges user medical profile + live AQI data; EVA Assistant prepares suggestion payload.",
        "Step 8 (UI Rendering): React frontend updates Leaflet spatial map pins, Recharts trend charts, and EVA assistant widgets."
    ], TEAL_PRIMARY)

    # -------------------------------------------------------------
    # SLIDE 10: 5. Conclusion and Discussion
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, BG_LIGHT)
    add_header(slide10, "5. Conclusion and Discussion", "TOPIC 5: CONCLUSION & DISCUSSION")

    add_card(slide10, Inches(0.8), Inches(1.65), Inches(5.65), Inches(5.1), "Key Technical Achievements", [
        "High Efficiency: Microservices design delivers sub-second response times with 85% cache hit efficiency.",
        "Effective Early Warnings: Isolation Forest accurately isolates pollution anomalies prior to widespread exposure.",
        "Clinical Value: Downloadable PDF health reports bridge the gap between environmental monitoring and patient care.",
        "Security Standard: Multi-tier RBAC, rate limiting, and security headers ensure robust data protection."
    ], TEAL_PRIMARY)

    add_card(slide10, Inches(6.85), Inches(1.65), Inches(5.65), Inches(5.1), "System Discussion & Future Scope", [
        "Current Trade-off: Dependent on external cloud API feeds; hyper-local precision varies with station density.",
        "Future IoT Integration: Connect ESP32 / Arduino hardware air sensors over MQTT for hyper-local microclimates.",
        "Deep Learning Models: Upgrade forecasting models from linear regression to LSTM / Transformer neural networks.",
        "Mobile App Expansion: Deploy React Native mobile app with background GPS push alerts for high-exposure zones."
    ], AMBER_GOLD)

    # -------------------------------------------------------------
    # SLIDE 11: 6. Screenshots & Other Diagrams
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, BG_LIGHT)
    add_header(slide11, "6. Screenshots & UI Demonstration Diagrams", "TOPIC 6: SCREENSHOTS & UI")

    add_card(slide11, Inches(0.8), Inches(1.65), Inches(3.65), Inches(5.1), "1. Main Dashboard & GIS Map", [
        "Interactive Leaflet spatial heatmap with station pins.",
        "Color-coded EPA AQI status pills (Good to Hazardous).",
        "6-Pollutant breakdown grid (PM2.5, PM10, NO2, etc.).",
        "Location search with auto-suggest."
    ], TEAL_PRIMARY)

    add_card(slide11, Inches(4.85), Inches(1.65), Inches(3.65), Inches(5.1), "2. EVA Assistant & Alerts", [
        "Menu-driven virtual assistant floating widget.",
        "Smart outdoor travel & itinerary recommendations.",
        "Dynamic AQI hazard alert banner notifications.",
        "Nearby healthcare clinic navigation."
    ], AMBER_GOLD)

    add_card(slide11, Inches(8.9), Inches(1.65), Inches(3.65), Inches(5.1), "3. Exposure & PDF Reports", [
        "Personalized cumulative exposure dose tracker.",
        "Vulnerability health profile setup screen.",
        "One-click downloadable PDF summary reports.",
        "Multi-tier RBAC admin management dashboards."
    ], SKY_BLUE)

    # -------------------------------------------------------------
    # SLIDE 12: Conclusion / Q&A (Dark Midnight Theme)
    # -------------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, BG_MIDNIGHT)

    # Decorative gold bar
    dec_bar12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.2), Inches(0.8), Inches(0.06))
    dec_bar12.fill.solid()
    dec_bar12.fill.fore_color.rgb = AMBER_GOLD
    dec_bar12.line.fill.background()

    tb12 = slide12.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.733), Inches(4.0))
    tf12 = tb12.text_frame
    tf12.word_wrap = True
    tf12.margin_left = tf12.margin_right = tf12.margin_top = tf12.margin_bottom = 0

    p12_1 = tf12.paragraphs[0]
    p12_1.text = "Thank You!"
    p12_1.font.size = Pt(40)
    p12_1.font.bold = True
    p12_1.font.color.rgb = TEXT_WHITE
    p12_1.space_after = Pt(14)

    p12_2 = tf12.add_paragraph()
    p12_2.text = "Environmental Health Intelligence & Early Warning System (EHI-EWS)"
    p12_2.font.size = Pt(20)
    p12_2.font.color.rgb = TEAL_PRIMARY
    p12_2.space_after = Pt(24)

    p12_3 = tf12.add_paragraph()
    p12_3.text = "Questions & Academic Discussion"
    p12_3.font.size = Pt(16)
    p12_3.font.bold = True
    p12_3.font.color.rgb = AMBER_GOLD

    # Save Presentation
    output_path = os.path.join("d:\\miniproject\\aqi p", "EHI_EWS_Elegant_Presentation.pptx")
    prs.save(output_path)
    print(f"Elegant Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    build_elegant_presentation()
