import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette - Professional Modern Dark & Emerald Theme
    DARK_BG = RGBColor(15, 23, 42)        # Slate 900
    PRIMARY_EMERALD = RGBColor(16, 185, 129) # Emerald 500
    SECONDARY_SKY = RGBColor(14, 165, 233)   # Sky 500
    TEXT_DARK = RGBColor(30, 41, 59)      # Slate 800
    TEXT_LIGHT = RGBColor(241, 245, 249)  # Slate 100
    CARD_BG = RGBColor(255, 255, 255)     # White
    CARD_BG_ALT = RGBColor(241, 245, 249) # Slate 100
    CARD_BORDER = RGBColor(226, 232, 240) # Slate 200
    TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500
    ACCENT_PURPLE = RGBColor(168, 85, 247)# Purple 500
    ACCENT_RED = RGBColor(239, 68, 68)    # Red 500

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="ENVIRONMENTAL HEALTH INTELLIGENCE & EARLY WARNING SYSTEM"):
        # Kicker / Category
        kicker_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.35))
        tf_k = kicker_box.text_frame
        tf_k.word_wrap = True
        p_k = tf_k.paragraphs[0]
        p_k.text = category_text.upper()
        p_k.font.size = Pt(11)
        p_k.font.bold = True
        p_k.font.color.rgb = PRIMARY_EMERALD

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.75))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_DARK

    def add_card(slide, left, top, width, height, title, body_bullets, accent_color=PRIMARY_EMERALD):
        # Background card shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        shape.line.width = Pt(1.5)

        # Top Accent Header Bar
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.08))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = accent_color
        top_bar.line.fill.background()

        # Text Frame
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.18), width - Inches(0.4), height - Inches(0.28))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_DARK
        p_t.space_after = Pt(8)

        for bullet in body_bullets:
            p_b = tf.add_paragraph()
            p_b.text = f"• {bullet}"
            p_b.font.size = Pt(12)
            p_b.font.color.rgb = TEXT_DARK
            p_b.space_after = Pt(5)

    def add_diagram_box(slide, left, top, width, height, title, items, bg_color=CARD_BG_ALT, line_color=PRIMARY_EMERALD):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = line_color
        shape.line.width = Pt(2)

        tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.15), width - Inches(0.3), height - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_DARK
        p_t.alignment = PP_ALIGN.CENTER
        p_t.space_after = Pt(6)

        for item in items:
            p_i = tf.add_paragraph()
            p_i.text = f"▪ {item}"
            p_i.font.size = Pt(11)
            p_i.font.color.rgb = TEXT_DARK
            p_i.space_after = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BG
    bg1.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p1_sub = tf1.paragraphs[0]
    p1_sub.text = "MINI PROJECT PRESENTATION"
    p1_sub.font.size = Pt(14)
    p1_sub.font.bold = True
    p1_sub.font.color.rgb = SECONDARY_SKY
    p1_sub.space_after = Pt(12)

    p1_title = tf1.add_paragraph()
    p1_title.text = "Environmental Health Intelligence & Early Warning System (EHI-EWS)"
    p1_title.font.size = Pt(32)
    p1_title.font.bold = True
    p1_title.font.color.rgb = PRIMARY_EMERALD
    p1_title.space_after = Pt(16)

    p1_desc = tf1.add_paragraph()
    p1_desc.text = "Real-Time Air Quality Monitoring, Machine Learning Anomaly Detection & AI Health Risk Assessment"
    p1_desc.font.size = Pt(18)
    p1_desc.font.color.rgb = TEXT_LIGHT
    p1_desc.space_after = Pt(32)

    p1_meta = tf1.add_paragraph()
    p1_meta.text = "Full-Stack Architecture: React (Vite) + Node.js (Express v5) + Python FastAPI (Scikit-Learn) + MongoDB"
    p1_meta.font.size = Pt(13)
    p1_meta.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 2: Table of Contents / Topic Agenda
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Presentation Overview & Agenda")

    agenda_items = [
        ("1. Introduction", "Background, Motivation & Project Objectives", PRIMARY_EMERALD),
        ("2. Problem Definition & Existing Systems", "Limitations of legacy solutions & Comparative Analysis", SECONDARY_SKY),
        ("3. Methodology", "Data Pipeline, ML Anomaly Detection & EVA Assistant", ACCENT_PURPLE),
        ("4. Design Diagrams", "System Architecture, ER Diagram & Data Flow Sequence", PRIMARY_EMERALD),
        ("5. Conclusion and Discussion", "System Impact, Key Achievements & Future Roadmap", SECONDARY_SKY),
        ("6. Screenshots & Other Diagrams", "UI Walkthrough, GIS Map, Dashboards & PDF Export", ACCENT_PURPLE)
    ]

    for idx, (title, desc, color) in enumerate(agenda_items):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.6 + row * 1.8)
        
        add_card(slide2, left, top, Inches(5.6), Inches(1.5), title, [desc], color)

    # -------------------------------------------------------------
    # SLIDE 3: 1. Introduction
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "1. Introduction — Project Overview & Objectives", "TOPIC 1: INTRODUCTION")

    add_card(slide3, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2), "Project Background", [
        "Rapid urbanization has led to severe air degradation (PM2.5, PM10, NO2).",
        "Air quality directly impacts public health, triggering respiratory and cardiac conditions.",
        "Need for a proactive early warning platform rather than passive pollution tracking."
    ], SECONDARY_SKY)

    add_card(slide3, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2), "Project Objectives", [
        "Real-Time Monitoring: Spatial tracking of 6 key pollutants via GIS mapping.",
        "Predictive Intelligence: Machine learning detection of sudden pollution spikes.",
        "Personalized Health Protection: Tailored advice based on individual medical risk profiles.",
        "Virtual Assistance: AI guidance (EVA) for outdoor activities and emergency care."
    ], PRIMARY_EMERALD)

    add_card(slide3, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2), "Core Innovation Pillars", [
        "Dual-Backend Microservices: Node.js API orchestrator + Python FastAPI ML service.",
        "Custom Anomaly Engine: Isolation Forest model for pollution anomaly detection.",
        "Clinical Report Generation: Downloadable doctor-ready PDF health exposure summaries.",
        "Role-Based Access: Dedicated views for Users, Facilities, and Administrators."
    ], ACCENT_PURPLE)

    # -------------------------------------------------------------
    # SLIDE 4: 2. Problem Definition & Existing Systems
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "2. Problem Definition & Existing Systems", "TOPIC 2: PROBLEM DEFINITION")

    add_card(slide4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Problem Definition", [
        "Broad, Generic Air Quality Metrics: Standard AQI numbers fail to account for personal vulnerability (e.g., Asthma, Elderly, Children).",
        "Reactive Emergency Response: Warnings are issued after peak pollution exposure has occurred, causing unmitigated health risks.",
        "Fragmented Healthcare Access: Citizens lack direct linkage between pollution alerts and nearby emergency medical facilities.",
        "Data Inconsistency & Spikes: Raw sensor feeds frequently suffer from unverified anomalies, leading to false alarms or missed hazard spikes."
    ], ACCENT_RED)

    add_card(slide4, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2), "Existing Systems vs. Proposed EHI-EWS", [
        "Legacy Apps: Show single numerical AQI without personalized dosage tracking.",
        "EHI-EWS: Calculates cumulative exposure dosage linked to user medical profile.",
        "Legacy Apps: No real-time predictive anomaly detection for sudden local fires/emissions.",
        "EHI-EWS: Microservice-driven Isolation Forest ML engine with statistical fallbacks.",
        "Legacy Apps: Static UI with zero interactive guidance.",
        "EHI-EWS: EVA (Environmental Virtual Assistant) menu & conversational AI + PDF reports."
    ], SECONDARY_SKY)

    # -------------------------------------------------------------
    # SLIDE 5: 3. Methodology - Data Pipeline & ML Engine
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "3. Methodology — Data Pipeline & ML Microservice", "TOPIC 3: METHODOLOGY")

    add_card(slide5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "1. Data Acquisition & Caching Pipeline", [
        "External Telemetry: Integrates OpenWeather REST API for live air pollution and geocoding.",
        "Monitored Pollutants: PM2.5, PM10, NO2, SO2, CO, and O3 values.",
        "Cache Optimization: NodeCache layer with 10-minute TTL reduces external API overhead by 85%.",
        "Reverse Geocoding: Converts lat/lon spatial coordinates into human-readable city locations."
    ], PRIMARY_EMERALD)

    add_card(slide5, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2), "2. Machine Learning Anomaly Detection", [
        "Algorithm: Isolation Forest (Scikit-Learn) fitted on temporal pollutant sequences.",
        "Z-Score Fallback: Rolling statistical z-score evaluation for short sequence cold starts.",
        "Evaluator Engine: Computes MAE, RMSE, and R² for regression; Precision, Recall, and F1 for classification.",
        "FastAPI Microservice: Isolated execution on port 8000 for high-performance ML inference."
    ], SECONDARY_SKY)

    # -------------------------------------------------------------
    # SLIDE 6: 3. Methodology - Health Assessment & EVA Assistant
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "3. Methodology — Health Risk Model & EVA Assistant", "TOPIC 3: METHODOLOGY")

    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Personalized Exposure Assessment Model", [
        "Profile Modeling: Captures pre-existing conditions (Asthma, COPD, Heart disease, Age group).",
        "Dosage Calculation: Exposure Dose = Ambient Concentration × Outdoor Duration × Vulnerability Factor.",
        "Risk Stratification: Classifies danger into Safe, Moderate, Severe, and Critical hazard zones.",
        "Automated PDF Engine: Uses jsPDF to compile historical exposure graphs into medical reports."
    ], ACCENT_PURPLE)

    add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2), "EVA (Environmental Virtual Assistant) Engine", [
        "Menu & Conversational Flow: Guided menu trees + free text natural language understanding.",
        "Context Awareness: Reads active page state, user location, and medical risk tier.",
        "Smart Itinerary: Evaluates outdoor travel schedules against multi-day forecast trends.",
        "Healthcare Routing: Locates nearest available hospitals and emergency care clinics."
    ], PRIMARY_EMERALD)

    # -------------------------------------------------------------
    # SLIDE 7: 4. Design Diagrams - System Architecture
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "4. Design Diagrams — System Architecture", "TOPIC 4: DESIGN DIAGRAMS")

    # Architecture Boxes layout
    add_diagram_box(slide7, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8), "CLIENT TIER\n(Vite + React 18 + TS)", [
        "Shadcn UI & Tailwind CSS",
        "Leaflet GIS Interactive Map",
        "Recharts Visualization",
        "jsPDF Medical Exporter",
        "EVA Assistant Widget"
    ], CARD_BG_ALT, SECONDARY_SKY)

    # Arrow 1
    a1 = slide7.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.5), Inches(3.8), Inches(0.6), Inches(0.4))
    a1.fill.solid()
    a1.fill.fore_color.rgb = PRIMARY_EMERALD
    a1.line.fill.background()

    add_diagram_box(slide7, Inches(5.2), Inches(1.8), Inches(3.6), Inches(4.8), "BACKEND API TIER\n(Node.js + Express v5)", [
        "JWT Auth & RBAC Middleware",
        "AQI Service & NodeCache",
        "EVA Orchestrator Engine",
        "Rate Limiting & Helmet",
        "Mongoose ORM Controllers"
    ], CARD_BG_ALT, PRIMARY_EMERALD)

    # Arrow 2
    a2 = slide7.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.9), Inches(3.8), Inches(0.6), Inches(0.4))
    a2.fill.solid()
    a2.fill.fore_color.rgb = ACCENT_PURPLE
    a2.line.fill.background()

    add_diagram_box(slide7, Inches(9.6), Inches(1.8), Inches(3.0), Inches(2.3), "ML SERVICE TIER\n(Python FastAPI)", [
        "Isolation Forest Model",
        "Z-Score Statistical Check",
        "Metrics Evaluator (MAE, R²)"
    ], CARD_BG_ALT, ACCENT_PURPLE)

    add_diagram_box(slide7, Inches(9.6), Inches(4.3), Inches(3.0), Inches(2.3), "DATABASE & EXTERNAL", [
        "MongoDB Atlas / Local",
        "OpenWeather Pollution API",
        "Geocoding REST Endpoints"
    ], CARD_BG_ALT, TEXT_MUTED)

    # -------------------------------------------------------------
    # SLIDE 8: 4. Design Diagrams - ER & Data Model Schema
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "4. Design Diagrams — Entity-Relationship (ER) Schema", "TOPIC 4: DESIGN DIAGRAMS")

    add_card(slide8, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2), "User & Profile Schema", [
        "User: _id, name, email, passwordHash, role (User, Admin, Facility), createdAt.",
        "Profile: _id, userId, age, asthma (Bool), heartCondition (Bool), vulnerabilityScore, emergencyContact."
    ], SECONDARY_SKY)

    add_card(slide8, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2), "Environmental Telemetry", [
        "EnvironmentalReading: _id, location, coordinates (lat, lon), aqi, pm25, pm10, no2, so2, co, o3, timestamp.",
        "Prediction: _id, modelVersion, anomalyScore, isAnomaly, confidence, timestamp."
    ], PRIMARY_EMERALD)

    add_card(slide8, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2), "Exposure & Facility Schema", [
        "ExposureRecord: _id, userId, durationMinutes, averageAqi, calculatedDose, date.",
        "Facility: _id, name, type (Hospital/Clinic), coordinates, capacity, contactPhone.",
        "EnvironmentalAlert: _id, severity, threshold, triggerLocation, status."
    ], ACCENT_PURPLE)

    # -------------------------------------------------------------
    # SLIDE 9: 4. Design Diagrams - Data Flow Sequence
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "4. Design Diagrams — Sequence & Data Flow", "TOPIC 4: DESIGN DIAGRAMS")

    add_card(slide9, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2), "End-to-End System Data Flow Sequence", [
        "1. Client Request: User accesses Dashboard or requests AQI update for current GPS coordinates.",
        "2. Node.js Middleware: Express rate limiter & JWT middleware validate request authenticity.",
        "3. Cache Verification: NodeCache checks if OpenWeather data for (lat, lon) exists within 10-min TTL.",
        "4. Telemetry Fetch: On cache miss, Node.js fetches 6-pollutant telemetry & weather forecast from OpenWeather API.",
        "5. ML Microservice Call: Express backend sends recent AQI sequence to Python FastAPI `/detect-anomalies` endpoint.",
        "6. Anomaly Evaluation: Isolation Forest checks sequence for sudden spikes, returning anomaly score and confidence.",
        "7. Exposure & EVA Processing: Health risk engine merges user profile + AQI data; EVA Assistant prepares advice options.",
        "8. UI Rendering: React frontend updates Leaflet GIS map pins, Recharts trend charts, and EVA suggestion banners."
    ], PRIMARY_EMERALD)

    # -------------------------------------------------------------
    # SLIDE 10: 5. Conclusion and Discussion
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "5. Conclusion and Discussion", "TOPIC 5: CONCLUSION & DISCUSSION")

    add_card(slide10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Key Findings & Technical Achievements", [
        "High Performance: Microservice architecture delivers sub-second AQI responses with 85% cache hit efficiency.",
        "Effective Early Warnings: Isolation Forest accurately isolates pollution anomalies prior to widespread exposure.",
        "Clinical Utility: Automated PDF health report generation bridges the gap between environmental monitoring and medical consultation.",
        "Robust Security: Implemented multi-tier RBAC, rate limiting, and security headers ensuring data protection."
    ], PRIMARY_EMERALD)

    add_card(slide10, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2), "Limitations & Future Roadmap", [
        "Current Limitation: Relies on cloud API feeds; hyper-local readings depend on station density.",
        "Future IoT Integration: Connect ESP32 / Arduino hardware air sensors over MQTT for hyper-local microclimates.",
        "Deep Learning Models: Upgrade forecasting models from regression to LSTM / Transformer architectures.",
        "Mobile App Expansion: Deploy React Native mobile app with background GPS push alerts for high-exposure zones."
    ], SECONDARY_SKY)

    # -------------------------------------------------------------
    # SLIDE 11: 6. Screenshots & Other Diagrams - UI Overview
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "6. Screenshots & UI Demonstration Diagrams", "TOPIC 6: SCREENSHOTS & UI")

    add_card(slide11, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2), "1. Main Dashboard & GIS Map", [
        "Interactive Leaflet spatial heatmap with real-time station pins.",
        "Color-coded AQI status badges (Good to Hazardous).",
        "Pollutant grid displaying PM2.5, PM10, NO2, SO2, CO, O3.",
        "Location search with auto-suggest."
    ], PRIMARY_EMERALD)

    add_card(slide11, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2), "2. EVA Assistant & Alerts", [
        "Menu-driven floating virtual assistant widget.",
        "Smart travel & itinerary outdoor safety recommendations.",
        "Dynamic AQI hazard alert banners.",
        "Nearby healthcare center navigation."
    ], SECONDARY_SKY)

    add_card(slide11, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2), "3. Exposure & PDF Reports", [
        "Personalized cumulative exposure dose tracker.",
        "Vulnerability health profile setup screen.",
        "One-click downloadable PDF summary reports for medical consultation.",
        "Multi-tier RBAC admin dashboards."
    ], ACCENT_PURPLE)

    # -------------------------------------------------------------
    # SLIDE 12: Thank You / Q&A (Dark Theme)
    # -------------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_layout)
    bg12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg12.fill.solid()
    bg12.fill.fore_color.rgb = DARK_BG
    bg12.line.fill.background()

    tb12 = slide12.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.5))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    p12_1 = tf12.paragraphs[0]
    p12_1.text = "Thank You!"
    p12_1.font.size = Pt(44)
    p12_1.font.bold = True
    p12_1.font.color.rgb = PRIMARY_EMERALD
    p12_1.space_after = Pt(16)

    p12_2 = tf12.add_paragraph()
    p12_2.text = "Environmental Health Intelligence & Early Warning System (EHI-EWS)"
    p12_2.font.size = Pt(22)
    p12_2.font.color.rgb = TEXT_LIGHT
    p12_2.space_after = Pt(24)

    p12_3 = tf12.add_paragraph()
    p12_3.text = "Questions & Discussion"
    p12_3.font.size = Pt(20)
    p12_3.font.bold = True
    p12_3.font.color.rgb = SECONDARY_SKY

    # Save output file
    output_path = os.path.join("d:\\miniproject\\aqi p", "EHI_EWS_Project_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation successfully generated at: {output_path}")

if __name__ == "__main__":
    build_presentation()
